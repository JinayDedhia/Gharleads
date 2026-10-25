"""GharLeads v0.4 - Mumbai real estate lead finder (Windows).  Run: python gharleads.py"""
import sys, os, re, json, time, asyncio, datetime, urllib.request, urllib.parse
if getattr(sys, "frozen", False):
    os.environ["PLAYWRIGHT_BROWSERS_PATH"] = "0"
from PySide6.QtCore import Qt, QThread, Signal, QPointF
from PySide6.QtGui import QPainter, QColor, QPen, QPainterPath, QIcon
from PySide6.QtWidgets import (QApplication, QWidget, QMainWindow, QStackedWidget, QVBoxLayout,
    QHBoxLayout, QGridLayout, QLabel, QLineEdit, QPushButton, QCheckBox, QComboBox, QSpinBox,
    QTabWidget, QGroupBox, QTableWidget, QTableWidgetItem, QProgressBar, QFileDialog,
    QMessageBox, QFrame, QHeaderView, QInputDialog)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- DATA
def S(region, names): return [(n.strip(), region) for n in names.split(",")]
LINES = {
 "Western Line": S("Mumbai", "Churchgate,Marine Lines,Charni Road,Grant Road,Mumbai Central,Mahalaxmi,Lower Parel,Prabhadevi,Dadar,Matunga Road,Mahim,Bandra,Khar Road,Santacruz,Vile Parle,Andheri,Jogeshwari,Ram Mandir,Goregaon,Malad,Kandivali,Borivali,Dahisar")
   + S("Thane", "Mira Road,Bhayandar") + S("Palghar", "Naigaon,Vasai Road,Nalla Sopara,Virar"),
 "Central Line": S("Mumbai", "CSMT,Masjid,Sandhurst Road,Byculla,Chinchpokli,Currey Road,Parel,Dadar,Matunga,Sion,Kurla,Vidyavihar,Ghatkopar,Vikhroli,Kanjurmarg,Bhandup,Nahur,Mulund")
   + S("Thane", "Thane,Kalva,Mumbra,Diva,Dombivli,Thakurli,Kalyan,Vithalwadi,Ulhasnagar,Ambernath,Badlapur,Shahad,Ambivli,Titwala,Vasind,Asangaon,Kasara")
   + S("Raigad", "Vangani,Neral,Karjat,Khopoli"),
 "Harbour Line": S("Mumbai", "CSMT,Masjid,Sandhurst Road,Dockyard Road,Reay Road,Cotton Green,Sewri,Wadala Road,Kings Circle,GTB Nagar,Chunabhatti,Kurla,Tilak Nagar,Chembur,Govandi,Mankhurd")
   + S("Navi Mumbai", "Vashi,Sanpada,Juinagar,Nerul,Seawoods Darave,CBD Belapur,Kharghar,Mansarovar,Khandeshwar,Panvel"),
 "Trans-Harbour": S("Thane", "Thane") + S("Navi Mumbai", "Airoli,Rabale,Ghansoli,Kopar Khairane,Turbhe,Sanpada,Juinagar,Vashi,Nerul"),
 "Uran Line": S("Navi Mumbai", "Nerul,Seawoods Darave,Sagar Sangam,Targhar,Bamandongri,Kharkopar,Ulwe,Gavhan,Ranjanpada,Nhava Sheva,Dronagiri,Uran"),
}  # best-effort station list - edit freely
LINE_OF = {}
for _l, _s in LINES.items():
    for _k in _s: LINE_OF.setdefault(_k, _l)
CATS = {
 "Real Estate Agents": ["real estate agent", "property agent"],
 "Real Estate Brokers": ["real estate broker", "property dealer"],
 "Channel Partners": ["real estate channel partner", "channel partner property"],
 "Regional Channel Partners": ["real estate consultant", "property consultant"],
 "Institutional Channel Partners": ["institutional channel partner real estate", "real estate advisory firm"],
}
ORDER = ["Business Name", "Phone", "Address", "Email", "Website", "Station", "Line", "Category", "Type", "Source", "WhatsApp", "Link"]
WIDTH = {"Business Name": 34, "Phone": 16, "Address": 52, "Email": 30, "Website": 34, "Station": 18, "Line": 15, "Category": 28, "Type": 24, "Source": 14, "WhatsApp": 34, "Link": 42}
INFO = {"Phone": ["Phone"], "Address": ["Address"], "Email (slower)": ["Email"], "Website": ["Website"], "Station / Line": ["Station", "Line"],
        "WhatsApp link": ["WhatsApp"], "Link to listing": ["Link"], "Business type": ["Type"]}
VERSION = "0.4"
OUT_DIR = os.path.join(os.path.expanduser("~"), "Documents", "GharLeads")
GEO_FILE = os.path.join(OUT_DIR, "station_coords.json")
SET_FILE = os.path.join(OUT_DIR, "settings.json")
MEM_FILE = os.path.join(OUT_DIR, "lead_memory.json")
REL_TYPE = re.compile(r"real estate|property|estate agen|realt|builder|develop|housing|apartment|flat|rental|leasing|land|plot|township|construction", re.I)
REL_NAME = re.compile(r"realt|propert|estate|housing|builder|develop|broker|dealer|channel|flat|homes|land|rental|infra|construction", re.I)
def jload(p, d):
    try: return json.load(open(p, encoding="utf-8"))
    except Exception: return d
def jsave(p, v):
    try: os.makedirs(OUT_DIR, exist_ok=True); json.dump(v, open(p, "w", encoding="utf-8"))
    except Exception: pass
def res(n): return os.path.join(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))), n)
def wa(phone): return f"https://wa.me/91{phone}" if re.fullmatch(r"[6-9]\d{9}", phone or "") else ""
DEEP = {"Real Estate Agents": ["estate agent", "rental agent"], "Real Estate Brokers": ["flat broker", "property broker"],
        "Channel Partners": ["property marketing partner"], "Regional Channel Partners": ["property consultants"],
        "Institutional Channel Partners": ["realty", "builders and developers"]}
OFFS = [(0, 0), (0.008, 0), (-0.008, 0), (0, 0.008), (0, -0.008)]   # map grid around a station (~900 m steps)
INST = re.compile(r"\b(pvt|private|ltd|limited|llp|developers?|realty|infra\w*)\b", re.I)

# ---------------------------------------------------------------- HELPERS
PH = re.compile(r"(?<!\d)(?:\+?91[\s-]*)?0?([6-9]\d{4}[\s-]?\d{5})(?!\d)|(?<!\d)(0?22[\s-]?\d{4}[\s-]?\d{4})(?!\d)")
def find_phone(t):
    m = PH.search(t or "")
    if not m: return ""
    d = re.sub(r"\D", "", m.group(1) or m.group(2))
    return d if m.group(1) else "022" + d[-8:]

def parse_addr(t):
    for l in (x.strip() for x in (t or "").split("\n")):
        if not re.search(r"[·⋅]", l) or re.match(r"(Open|Closed|Closes|Opens)", l): continue
        parts = [p.strip() for p in re.split(r"[·⋅]", l)]
        c = [p for p in parts[1:] if p and not find_phone(p) and not re.match(r"(Open|Closed|Closes|Opens)", p)]
        if c: return c[0]
    return ""

def parse_type(t):
    for l in (x.strip() for x in (t or "").split("\n")):
        if not re.search(r"[·⋅]", l) or re.match(r"(Open|Closed|Closes|Opens)", l): continue
        for p in (x.strip() for x in re.split(r"[·⋅]", l)):
            if p and not re.match(r"\d", p): return p if len(p) <= 40 and not re.search(r"\d", p) else ""
    return ""

def find_email(url):
    try:
        h = urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=7).read(400000).decode("utf-8", "ignore")
    except Exception: return ""
    m = re.findall(r"mailto:([\w.+-]+@[\w-]+\.[\w.-]+)", h) or re.findall(r"[\w.+-]+@[\w-]+\.[a-zA-Z]{2,}", h)
    bad = (".png", ".jpg", ".svg", ".gif", "sentry", "example", "wixpress")
    return next((e for e in m if not any(b in e.lower() for b in bad)), "")

def find_email_text(t):
    m = re.search(r"[\w.+-]+@[\w-]+\.[a-zA-Z]{2,}", t or "")
    return m.group(0) if m else ""

def http_json(url, data=None, timeout=30):
    req = urllib.request.Request(url, data=data, headers={"User-Agent": "GharLeads/0.3 (desktop lead finder)"})
    return json.loads(urllib.request.urlopen(req, timeout=timeout).read().decode("utf-8"))

def geocode(name, region):
    for q in (f"{name} railway station, {region}, Maharashtra, India", f"{name}, {region}, Maharashtra, India"):
        try: r = http_json("https://nominatim.openstreetmap.org/search?format=json&limit=1&q=" + urllib.parse.quote(q))
        except Exception: r = []
        time.sleep(1.1)   # Nominatim allows 1 request per second
        if r:
            la, lo = float(r[0]["lat"]), float(r[0]["lon"])
            if 18.5 < la < 20.2 and 72.6 < lo < 73.6: return la, lo
    return None

def osm_leads(lat, lon, r=2500):
    q = ('[out:json][timeout:25];(nwr["office"="estate_agent"](around:%d,%f,%f);nwr["shop"="estate_agent"](around:%d,%f,%f););out center tags;' % (r, lat, lon, r, lat, lon))
    return http_json("https://overpass-api.de/api/interpreter", data=urllib.parse.urlencode({"data": q}).encode(), timeout=45).get("elements", [])

def plan(c, geo):
    """All Google Maps searches for a run: normal ones first, deep-mode (map grid + extra keywords) after."""
    jobs = [(st, t, cat, None) for st in c["stations"] for cat in c["cats"] for t in CATS[cat]]
    if c["deep"]:
        for st in c["stations"]:
            ll = geo.get(st)
            if not ll: continue
            for da, do in OFFS:
                for cat in c["cats"]: jobs.append((st, CATS[cat][0], cat, (ll[0] + da, ll[1] + do)))
            for cat in c["cats"]:
                for t in DEEP[cat]: jobs.append((st, t, cat, ll))
    return jobs

def save_excel(path, rows, fields):
    cols = [f for f in ORDER if f in fields]
    wb = Workbook(); ws = wb.active; ws.title = "All Leads"
    def fill(w, data):
        w.append(cols)
        for r in data:
            w.append([r.get(c, "") for c in cols])
            for lc in ("Link", "WhatsApp"):
                if lc in cols and r.get(lc):
                    k = w.cell(w.max_row, cols.index(lc) + 1); k.hyperlink = r[lc]; k.font = Font(color="0563C1", underline="single")
        for i, c in enumerate(cols, 1):
            h = w.cell(1, i); h.font = Font(bold=True, color="FFFFFF"); h.fill = PatternFill("solid", start_color="3B5998")
            w.column_dimensions[get_column_letter(i)].width = WIDTH[c]
        w.freeze_panes = "A2"; w.auto_filter.ref = w.dimensions
    fill(ws, rows)
    for st in sorted({r["Station"] for r in rows}):
        fill(wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", st)[:31]), [r for r in rows if r["Station"] == st])
    wb.save(path)

JS = "els=>els.map(a=>{const c=a.closest('div.Nv2PK')||a.parentElement;return {n:a.getAttribute('aria-label'),h:a.href,t:c?c.innerText:''}})"

# ---------------------------------------------------------------- ENGINE

class Worker(QThread):
    row = Signal(dict); status = Signal(str); prog = Signal(int, int)
    def __init__(s, cfg): super().__init__(); s.cfg = cfg; s.stop = False; s.errors = 0
    def log(s, m):
        try:
            os.makedirs(OUT_DIR, exist_ok=True)
            open(os.path.join(OUT_DIR, "gharleads_log.txt"), "a", encoding="utf-8").write(f"{datetime.datetime.now():%d-%b %H:%M:%S}  {m}\n")
        except Exception: pass
    def run(s):
        try: asyncio.run(s.go())
        except Exception as e: s.log(f"FATAL {e!r}"); s.status.emit(f"Stopped: {str(e)[:120]}")
    async def reset(s):
        async with s.lock:
            if s.br and s.br.is_connected(): return
            s.br = await s.p.chromium.launch(headless=True, args=["--disable-blink-features=AutomationControlled", "--lang=en-IN"])
            s.ctx = await s.br.new_context(locale="en-IN", timezone_id="Asia/Kolkata", viewport={"width": 1280, "height": 900},
                                           geolocation={"latitude": 19.07, "longitude": 72.87}, permissions=["geolocation"])
            await s.ctx.route("**/*", s.block)
    async def block(s, route):
        if route.request.resource_type in ("image", "media", "font"): await route.abort()
        else: await route.continue_()
    def locate(s):
        try: cache = json.load(open(GEO_FILE, encoding="utf-8"))
        except Exception: cache = {}
        for st in s.cfg["stations"]:
            k = f"{st[0]}|{st[1]}"
            if k not in cache and not s.stop:
                s.status.emit(f"Locating {st[0]} on the map (first time only) ...")
                ll = geocode(*st)
                if ll:
                    cache[k] = ll
                    try: os.makedirs(OUT_DIR, exist_ok=True); json.dump(cache, open(GEO_FILE, "w", encoding="utf-8"))
                    except Exception: pass
                else: s.log(f"could not locate station {st[0]}")
        s.geo = {st: tuple(cache[f"{st[0]}|{st[1]}"]) for st in s.cfg["stations"] if f"{st[0]}|{st[1]}" in cache}
    async def go(s):
        from playwright.async_api import async_playwright
        c = s.cfg; s.geo = {}
        if c["deep"] or c["osm"]: await asyncio.to_thread(s.locate)
        jobs = plan(c, s.geo); osm_st = [st for st in c["stations"] if st in s.geo] if c["osm"] else []
        s.total, s.n, s.seen, s.urls, s.count = len(jobs) + len(osm_st), 0, set(c["known"]), set(), {}
        s.sem, s.dsem, s.osem, s.lock, s.br, s.ctx = asyncio.Semaphore(c["tabs"]), asyncio.Semaphore(5), asyncio.Semaphore(1), asyncio.Lock(), None, None
        s.log(f"RUN START: {len(c['stations'])} stations, {len(c['cats'])} categories, max {c['max']}, tabs {c['tabs']}, deep={c['deep']}, osm={c['osm']}, {len(jobs)} map searches")
        async with async_playwright() as s.p:
            await s.reset()
            await asyncio.gather(*(s.osm(st) for st in osm_st), *(s.job(*j) for j in jobs))
            if s.br: await s.br.close()
        s.log(f"RUN END: {sum(s.count.values())} leads, {s.errors} errors")
    def cat_of(s, nm):
        n = nm.lower(); cats = s.cfg["cats"]
        g = ("Institutional Channel Partners" if INST.search(nm) or "builder" in n else "Regional Channel Partners" if re.search("consult|advis", n)
             else "Real Estate Brokers" if re.search("broker|dealer", n) else "Channel Partners" if "channel" in n else "Real Estate Agents")
        return g if g in cats else cats[0]
    async def osm(s, st):
        c = s.cfg; name, region = st; els = []
        async with s.osem:
            try:
                if not (s.stop or s.count.get(st, 0) >= c["max"]):
                    s.status.emit(f"{name}: OpenStreetMap ...")
                    els = await asyncio.to_thread(osm_leads, *s.geo[st], c["osm_r"])
                    await asyncio.sleep(1.5)
            except Exception as e:
                s.errors += 1; s.log(f"OSM {name}: {e!r}"[:250])
            finally:
                s.n += 1; s.prog.emit(s.n, s.total)
        for el in els:
            t = el.get("tags", {}); nm = (t.get("name") or "").strip()
            if not nm: continue
            phone = find_phone(" ".join(str(t.get(k, "")) for k in ("phone", "contact:phone", "mobile", "contact:mobile")))
            if c["phone_only"] and not phone: continue
            addr = t.get("addr:full") or ", ".join(t[k] for k in ("addr:housenumber", "addr:street", "addr:suburb", "addr:city") if t.get(k))
            k = phone or (nm + addr[:30]).lower()
            if k in s.seen or s.count.get(st, 0) >= c["max"]: continue
            s.seen.add(k); s.count[st] = s.count.get(st, 0) + 1
            s.row.emit({"Business Name": nm, "Phone": phone, "Address": addr or f"Near {name}, {region}", "Email": t.get("email") or t.get("contact:email") or "",
                        "Website": t.get("website") or t.get("contact:website") or "", "Station": name, "Line": LINE_OF.get(st, ""),
                        "Category": s.cat_of(nm), "Type": "", "Source": "OpenStreetMap", "WhatsApp": wa(phone), "_key": k, "Link": f"https://www.openstreetmap.org/{el['type']}/{el['id']}"})
    async def job(s, st, term, cat, ll):
        async with s.sem:
            if not s.stop and s.count.get(st, 0) < s.cfg["max"]:   # station already full -> skip remaining searches
                for attempt in (1, 2):
                    if s.stop: break
                    try:
                        s.status.emit(f"{st[0]}: searching '{term}' ...")
                        await s.search(st, term, cat, ll); break
                    except Exception as e:
                        s.errors += 1; s.log(f"{st[0]} / {term} (try {attempt}): {e!r}"[:300])
                        if not (s.br and s.br.is_connected()): await s.reset()
                        await asyncio.sleep(3)
            s.n += 1; s.prog.emit(s.n, s.total)
    async def search(s, st, term, cat, ll=None):
        c = s.cfg; name, region = st
        page = await s.ctx.new_page()
        try:
            if ll: url = f"https://www.google.com/maps/search/{term.replace(' ', '+')}/@{ll[0]:.5f},{ll[1]:.5f},16z?hl=en"
            else: url = "https://www.google.com/maps/search/" + f"{term} near {name} railway station {region}".replace(" ", "+") + "?hl=en"
            await page.goto(url, timeout=30000, wait_until="domcontentloaded")
            if "/sorry" in page.url:
                await asyncio.sleep(30); raise RuntimeError("Google asked for verification - slowing down")
            for b in ("Accept all", "Reject all"):
                try: await page.click(f"button:has-text('{b}')", timeout=1000)
                except Exception: pass
            try: await page.wait_for_selector("div[role='feed']", timeout=15000)
            except Exception: s.log(f"{name} / {term}: no results list (skipped)"); return
            while not s.stop:
                n = await page.locator("a.hfpxzc").count()
                if n >= c["max"]: break
                await page.eval_on_selector("div[role='feed']", "e=>e.scrollTo(0,e.scrollHeight)")
                try: await page.wait_for_function(f"document.querySelectorAll('a.hfpxzc').length>{n}", timeout=3500)
                except Exception: break
            items = await page.eval_on_selector_all("a.hfpxzc", JS)
        finally:
            await page.close()
        todo = [it for it in items[:c["max"]] if (it["n"] or "").strip() and it["h"].split("?")[0] not in s.urls]
        await asyncio.gather(*(s.item(it, st, cat) for it in todo))
    async def item(s, it, st, cat):
        c = s.cfg; name, region = st; url = it["h"].split("?")[0]
        if s.stop or s.count.get(st, 0) >= c["max"] or url in s.urls: return
        s.urls.add(url)
        try:
            nm = it["n"].strip(); phone, addr, web, typ = find_phone(it["t"]), parse_addr(it["t"]), "", parse_type(it["t"])
            if c["relevant"] and typ and not (REL_TYPE.search(typ) or REL_NAME.search(nm)): return   # unrelated business
            if (not phone and not c["fast"]) or c["web"]:
                async with s.dsem: d = await s.detail(it["h"])
                phone, addr, web = phone or d.get("phone", ""), addr or d.get("addr", ""), d.get("web", "")
            if c["phone_only"] and not phone: return
            k = phone or (nm + addr[:30]).lower()
            if k in s.seen or s.count.get(st, 0) >= c["max"]: return
            s.seen.add(k); s.count[st] = s.count.get(st, 0) + 1
            email = await asyncio.to_thread(find_email, web) if c["email"] and web else ""
            s.row.emit({"Business Name": nm, "Phone": phone, "Address": addr or f"Near {name}, {region}", "Email": email, "Website": web,
                        "Station": name, "Line": LINE_OF.get(st, ""), "Category": cat, "Type": typ, "Source": "Google Maps", "WhatsApp": wa(phone), "Link": url, "_key": k})
        except Exception as e:
            s.urls.discard(url); s.errors += 1; s.log(f"item error {e!r}"[:200])
    async def detail(s, url):
        d = {}; p = await s.ctx.new_page()
        try:
            await p.goto(url, timeout=25000, wait_until="domcontentloaded")
            await p.wait_for_selector("h1", timeout=8000)
            for k, sel, a in (("phone", "button[data-item-id^='phone']", "aria-label"),
                              ("addr", "button[data-item-id='address']", "aria-label"),
                              ("web", "a[data-item-id='authority']", "href")):
                try: v = await p.get_attribute(sel, a, timeout=1500) or ""
                except Exception: continue
                if k == "phone": d[k] = find_phone(v)
                elif k == "addr": d[k] = re.sub(r"^Address:\s*", "", v).strip()
                else: d[k] = v
        except Exception: pass
        finally: await p.close()
        return d

# ---------------------------------------------------------------- UI
QSS = """
*{font-family:Tahoma,Verdana;font-size:11px;color:#333}
QMainWindow,#page{background:#fff}
#home{background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #fff,stop:1 #d4dae8)}
#hdr{background:#3b5998;border-bottom:1px solid #29447e}
#hdr QLabel{color:#fff;font-weight:bold}
#logo{color:#fff;font:bold 34px Verdana}
#strip{background:#eceff5;border-bottom:1px solid #d8dfea}
#bar{background:#eceff5;border-top:1px solid #d8dfea}
#tag{font:bold 20px Tahoma;color:#1d2a57}
#h2{font:bold 20px Tahoma;color:#1d2a57}
QLineEdit{border:1px solid #96a6c5;background:#fff;padding:2px;min-height:18px}
QPushButton{background:#f7f7f7;border:1px solid #999;padding:4px 12px;font-weight:bold}
QPushButton#blue{background:#5b74a8;color:#fff;border:1px solid #29447e}
QPushButton#green{background:#69a74e;color:#fff;border:1px solid #3b6e22;font-size:13px;padding:6px 18px}
QPushButton#green:disabled{background:#a9c99b}
QGroupBox{border:1px solid #d8dfea;margin-top:16px;padding-top:8px}
QGroupBox::title{subcontrol-origin:margin;left:0;padding:4px 10px;background:#eceff5;color:#3b5998;font-weight:bold;font-size:12px}
QTabBar::tab{background:#eceff5;border:1px solid #d8dfea;padding:5px 12px;margin-right:3px}
QTabBar::tab:selected{background:#fff;font-weight:bold;color:#1d2a57}
QTabWidget::pane{border:1px solid #d8dfea}
QHeaderView::section{background:#3b5998;color:#fff;font-weight:bold;padding:4px;border:0}
QProgressBar{border:1px solid #96a6c5;text-align:center;height:16px}
QProgressBar::chunk{background:#69a74e}
"""

class Net(QWidget):
    PTS = [(30, 40), (130, 20), (220, 60), (310, 30), (400, 50), (50, 140), (130, 180), (230, 140), (320, 180), (400, 120)]
    def __init__(s): super().__init__(); s.setMinimumSize(440, 220)
    def paintEvent(s, e):
        p = QPainter(s); p.setRenderHint(QPainter.Antialiasing)
        p.setPen(QPen(QColor("#b5bfd8"), 3, Qt.DotLine))
        for i, (a, b) in enumerate(zip(s.PTS, s.PTS[1:] + s.PTS[:1])): p.drawLine(QPointF(*a), QPointF(*b))
        for a, b in ((1, 6), (3, 8), (2, 7), (0, 5)): p.drawLine(QPointF(*s.PTS[a]), QPointF(*s.PTS[b]))
        p.setPen(QPen(QColor("#c9771a"), 1.5)); p.setBrush(QColor("#f3a935"))
        for x, y in s.PTS:
            h = QPainterPath(); h.moveTo(x - 12, y); h.lineTo(x, y - 12); h.lineTo(x + 12, y)
            h.lineTo(x + 12, y + 12); h.lineTo(x - 12, y + 12); h.closeSubpath(); p.drawPath(h)

def header(right):
    f = QFrame(); f.setObjectName("hdr"); f.setFixedHeight(64)
    l = QHBoxLayout(f); l.setContentsMargins(24, 0, 24, 0)
    lg = QLabel("gharleads"); lg.setObjectName("logo"); l.addWidget(lg); l.addStretch(); l.addWidget(right)
    return f

def labelled(text, w):
    c = QWidget(); l = QVBoxLayout(c); l.setContentsMargins(0, 0, 0, 0); l.setSpacing(2); l.addWidget(QLabel(text)); l.addWidget(w); return c

class App(QMainWindow):
    def __init__(s):
        super().__init__(); s.setWindowTitle(f"GharLeads {VERSION}"); s.resize(1100, 900); s.setStyleSheet(QSS)
        if os.path.exists(res("icon.ico")): s.setWindowIcon(QIcon(res("icon.ico")))
        s.stack = QStackedWidget(); s.setCentralWidget(s.stack)
        s.stack.addWidget(s.home_page()); s.stack.addWidget(s.picker_page()); s.stack.addWidget(s.results_page())
        s.rows, s.worker = [], None
        s.load_presets(); last = jload(SET_FILE, {}).get("last")
        if last: s.apply(last)

    # ---- page 0: home (Facebook 2010 style)
    def home_page(s):
        pg = QWidget(); pg.setObjectName("home"); v = QVBoxLayout(pg); v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)
        r = QWidget(); rl = QHBoxLayout(r); rl.setContentsMargins(0, 0, 0, 0)
        em, lk = QLineEdit(), QLineEdit(); lk.setEchoMode(QLineEdit.Password); em.setFixedWidth(150); lk.setFixedWidth(150)
        go = QPushButton("Login"); go.setObjectName("blue"); go.clicked.connect(lambda: s.stack.setCurrentIndex(1))
        rl.addWidget(labelled("Email", em)); rl.addWidget(labelled("License Key", lk)); rl.addWidget(go, 0, Qt.AlignBottom)
        v.addWidget(header(r))
        body = QHBoxLayout(); body.setContentsMargins(48, 40, 48, 20); body.setSpacing(40)
        left = QVBoxLayout(); t = QLabel("GharLeads helps you find and connect with the real estate professionals in your city.")
        t.setObjectName("tag"); t.setWordWrap(True); left.addWidget(t); left.addSpacing(20); left.addWidget(Net()); left.addStretch()
        right = QVBoxLayout(); h = QLabel("Get started"); h.setObjectName("h2"); right.addWidget(h)
        sub = QLabel("Pick your stations, choose who you want to find, and export names, phone numbers and addresses to Excel.")
        sub.setWordWrap(True); right.addWidget(sub); right.addSpacing(14)
        st = QPushButton("Start Finding Leads"); st.setObjectName("green"); st.clicked.connect(lambda: s.stack.setCurrentIndex(1)); right.addWidget(st, 0, Qt.AlignLeft)
        right.addSpacing(14); right.addWidget(QLabel("Development build - licence check is not active yet.")); right.addStretch()
        body.addLayout(left, 3); body.addLayout(right, 2); v.addLayout(body, 1)
        ft = QLabel("Western Line   Central Line   Harbour Line   Trans-Harbour   Uran Line        GharLeads v" + VERSION + " (c) 2026"); ft.setContentsMargins(24, 8, 24, 8)
        ft.setStyleSheet("background:#fff"); v.addWidget(ft)
        return pg

    # ---- page 1: station picker
    def picker_page(s):
        pg = QWidget(); pg.setObjectName("page"); v = QVBoxLayout(pg); v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)
        v.addWidget(header(QLabel("Find Leads")))
        sp = QLabel("  Development build"); sp.setObjectName("strip"); sp.setFixedHeight(26); v.addWidget(sp)
        body = QHBoxLayout(); body.setContentsMargins(22, 6, 22, 10); body.setSpacing(18)
        # left: stations
        gb = QGroupBox("1. Pick your stations"); gl = QVBoxLayout(gb)
        top = QHBoxLayout(); top.addWidget(QLabel("Search:")); s.q = QLineEdit(); s.q.setPlaceholderText("e.g. Andheri"); s.q.setMaximumWidth(180)
        s.q.textChanged.connect(s.filter); top.addWidget(s.q)
        for txt, fn in (("Select line", lambda: s.setline(True)), ("Clear line", lambda: s.setline(False)), ("Clear all", s.clearall)):
            b = QPushButton(txt); b.clicked.connect(fn); top.addWidget(b)
        top.addStretch(); gl.addLayout(top)
        s.tabs = QTabWidget(); s.boxes = {}; s.line_names = list(LINES)
        for ln, sts in LINES.items():
            w = QWidget(); g = QGridLayout(w)
            for i, k in enumerate(sts):
                cb = QCheckBox(k[0]); cb.setProperty("nm", k[0].lower()); s.boxes.setdefault(k, []).append(cb)
                cb.toggled.connect(lambda val, k=k: s.sync(k, val)); g.addWidget(cb, i % 12, i // 12)
            g.setRowStretch(12, 1); s.tabs.addTab(w, ln)
        gl.addWidget(s.tabs); gl.addWidget(QLabel("Interchange stations (e.g. Dadar, Thane) are ticked on every line they serve."))
        body.addWidget(gb, 17)
        # right: categories, info, options
        col = QVBoxLayout()
        g2 = QGroupBox("2. Who are you looking for?"); l2 = QVBoxLayout(g2); s.cat = {}
        for c in CATS: cb = QCheckBox(c); cb.setChecked(True); s.cat[c] = cb; l2.addWidget(cb)
        g3 = QGroupBox("3. What info do you want?"); l3 = QVBoxLayout(g3)
        nm = QCheckBox("Business name (always on)"); nm.setChecked(True); nm.setEnabled(False); l3.addWidget(nm); s.info = {}
        for c, on in (("Phone", True), ("Address", True), ("Email (slower)", False), ("Website", False), ("Station / Line", True), ("WhatsApp link", True), ("Link to listing", True), ("Business type", False)):
            cb = QCheckBox(c); cb.setChecked(on); s.info[c] = cb; l3.addWidget(cb)
        g4 = QGroupBox("4. Options"); l4 = QVBoxLayout(g4)
        s.phone_only = QCheckBox("Only keep contacts that have a phone number"); s.phone_only.setChecked(True); l4.addWidget(s.phone_only)
        s.osm = QCheckBox("Also use OpenStreetMap (free, extra leads)"); s.osm.setChecked(True); l4.addWidget(s.osm)
        s.deep = QCheckBox("Deep mode (searches around each station - slower, finds more)"); l4.addWidget(s.deep)
        s.relevant = QCheckBox("Remove unrelated businesses (recommended)"); s.relevant.setChecked(True); l4.addWidget(s.relevant)
        s.only_new = QCheckBox("Only new leads (skip ones I already exported)"); s.only_new.setChecked(True); l4.addWidget(s.only_new)
        fb = QPushButton("Forget exported leads"); fb.clicked.connect(s.forget); l4.addWidget(fb, 0, Qt.AlignLeft)
        r = QHBoxLayout(); r.addWidget(QLabel("Max results per station:")); s.mx = QSpinBox(); s.mx.setRange(5, 500); s.mx.setValue(40); r.addWidget(s.mx); r.addStretch(); l4.addLayout(r)
        r = QHBoxLayout(); r.addWidget(QLabel("Speed:")); s.speed = QComboBox(); s.speed.addItems(["Normal (opens listings missing a phone)", "Fast (search list only)"]); r.addWidget(s.speed); l4.addLayout(r)
        for g in (g2, g3, g4): col.addWidget(g)
        col.addStretch(); body.addLayout(col, 10); v.addLayout(body, 1)
        bar = QFrame(); bar.setObjectName("bar"); bl = QHBoxLayout(bar); bl.setContentsMargins(22, 10, 22, 10)
        s.summary = QLabel(); bl.addWidget(s.summary); bl.addStretch()
        s.presets = QComboBox(); s.presets.setMinimumWidth(150); s.presets.activated.connect(s.use_preset); bl.addWidget(s.presets)
        for txt, fn in (("Save preset", s.save_preset), ("Delete preset", s.del_preset)):
            b = QPushButton(txt); b.clicked.connect(fn); bl.addWidget(b)
        s.start = QPushButton("Start Finding Leads"); s.start.setObjectName("green"); s.start.clicked.connect(s.begin); bl.addWidget(s.start)
        v.addWidget(bar); s.refresh(); return pg

    # ---- remembered settings, presets, lead memory
    def collect(s):
        return {"stations": [f"{k[0]}|{k[1]}" for k in s.picked()], "cats": [c for c, cb in s.cat.items() if cb.isChecked()],
                "info": {c: cb.isChecked() for c, cb in s.info.items()}, "max": s.mx.value(), "speed": s.speed.currentIndex(),
                "opts": {n: cb.isChecked() for n, cb in s.optboxes().items()}}
    def optboxes(s): return {"phone_only": s.phone_only, "osm": s.osm, "deep": s.deep, "relevant": s.relevant, "only_new": s.only_new}
    def apply(s, d):
        want = set(d.get("stations", []))
        for k, bs in s.boxes.items(): bs[0].setChecked(f"{k[0]}|{k[1]}" in want)
        for c, cb in s.cat.items(): cb.setChecked(c in d.get("cats", list(s.cat)))
        for c, cb in s.info.items(): cb.setChecked(d.get("info", {}).get(c, cb.isChecked()))
        for n, cb in s.optboxes().items(): cb.setChecked(d.get("opts", {}).get(n, cb.isChecked()))
        s.mx.setValue(d.get("max", s.mx.value())); s.speed.setCurrentIndex(d.get("speed", 0)); s.refresh()
    def load_presets(s):
        s.presets.clear(); s.presets.addItem("Presets ...")
        s.presets.addItems(sorted(jload(SET_FILE, {}).get("presets", {})))
    def use_preset(s, i):
        d = jload(SET_FILE, {}).get("presets", {}).get(s.presets.currentText())
        if i > 0 and d: s.apply(d)
    def save_preset(s):
        n, ok = QInputDialog.getText(s, "Save preset", "Preset name (e.g. My Western Line run):")
        if ok and n.strip():
            st = jload(SET_FILE, {}); st.setdefault("presets", {})[n.strip()] = s.collect(); jsave(SET_FILE, st)
            s.load_presets(); s.presets.setCurrentText(n.strip())
    def del_preset(s):
        n = s.presets.currentText(); st = jload(SET_FILE, {})
        if s.presets.currentIndex() > 0 and n in st.get("presets", {}): del st["presets"][n]; jsave(SET_FILE, st); s.load_presets()
    def forget(s):
        if QMessageBox.question(s, "GharLeads", "Forget all previously exported leads?\nThe next run will include them again.") == QMessageBox.Yes:
            jsave(MEM_FILE, []); QMessageBox.information(s, "GharLeads", "Done - exported leads forgotten.")

    def sync(s, k, val):
        for b in s.boxes[k]:
            if b.isChecked() != val: b.blockSignals(True); b.setChecked(val); b.blockSignals(False)
        s.refresh()
    def picked(s): return [k for k, bs in s.boxes.items() if bs[0].isChecked()]
    def refresh(s):
        sel = set(s.picked())
        for i, ln in enumerate(LINES):
            n = len([k for k in LINES[ln] if k in sel]); s.tabs.setTabText(i, f"{ln} ({n})" if n else ln)
        s.summary.setText(f"<b>{len(sel)}</b> stations selected")
        s.start.setEnabled(bool(sel))
    def setline(s, val):
        for k in LINES[s.line_names[s.tabs.currentIndex()]]: s.boxes[k][0].setChecked(val)
    def clearall(s):
        for bs in s.boxes.values(): bs[0].setChecked(False)
    def filter(s, t):
        t = t.lower().strip()
        for bs in s.boxes.values():
            for b in bs: b.setVisible(t in b.property("nm"))

    # ---- page 2: live results
    def results_page(s):
        pg = QWidget(); pg.setObjectName("page"); v = QVBoxLayout(pg); v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)
        v.addWidget(header(QLabel("Results")))
        body = QVBoxLayout(); body.setContentsMargins(22, 12, 22, 12)
        s.status = QLabel("Ready"); body.addWidget(s.status); s.pb = QProgressBar(); body.addWidget(s.pb)
        s.table = QTableWidget(); s.table.setEditTriggers(QTableWidget.NoEditTriggers); s.table.setAlternatingRowColors(True)
        s.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive); body.addWidget(s.table, 1)
        r = QHBoxLayout(); s.back = QPushButton("Back"); s.back.clicked.connect(lambda: s.stack.setCurrentIndex(1))
        s.stopb = QPushButton("Stop"); s.stopb.clicked.connect(s.halt)
        s.saveb = QPushButton("Save to Excel..."); s.saveb.setObjectName("green"); s.saveb.clicked.connect(s.save_as)
        for b in (s.back, s.stopb): r.addWidget(b)
        r.addStretch(); r.addWidget(s.saveb); body.addLayout(r); v.addLayout(body, 1); return pg

    def begin(s):
        cats = [c for c, cb in s.cat.items() if cb.isChecked()]
        if not cats: return QMessageBox.information(s, "GharLeads", "Please tick at least one category.")
        f = {"Business Name", "Category", "Source"}
        for c, cb in s.info.items():
            if cb.isChecked(): f |= set(INFO[c])
        st = jload(SET_FILE, {}); st["last"] = s.collect(); jsave(SET_FILE, st)
        s.fields = f; s.cols = [c for c in ORDER if c in f]; s.rows = []
        s.table.clear(); s.table.setRowCount(0); s.table.setColumnCount(len(s.cols)); s.table.setHorizontalHeaderLabels(s.cols)
        for i, c in enumerate(s.cols): s.table.setColumnWidth(i, WIDTH[c] * 7)
        cfg = dict(stations=s.picked(), cats=cats, max=s.mx.value(), phone_only=s.phone_only.isChecked(),
                   fast=s.speed.currentIndex() == 1, tabs=6 if s.speed.currentIndex() == 1 else 4, relevant=s.relevant.isChecked(),
                   known=jload(MEM_FILE, []) if s.only_new.isChecked() else [], deep=s.deep.isChecked(), osm=s.osm.isChecked(), osm_r=2500, email="Email" in f, web=bool({"Email", "Website"} & f))
        s.pb.setValue(0); s.stopb.setEnabled(True); s.back.setEnabled(False)
        s.autopath = os.path.join(OUT_DIR, "gharleads_" + datetime.datetime.now().strftime("%Y%m%d_%H%M") + ".xlsx")
        s.worker = Worker(cfg); s.worker.row.connect(s.add); s.worker.status.connect(s.status.setText)
        s.worker.prog.connect(lambda n, t: (s.pb.setMaximum(t), s.pb.setValue(n))); s.worker.finished.connect(s.done)
        s.worker.start(); s.stack.setCurrentIndex(2)

    def add(s, r):
        s.rows.append(r); i = s.table.rowCount(); s.table.insertRow(i)
        for j, c in enumerate(s.cols): s.table.setItem(i, j, QTableWidgetItem(r.get(c, "")))
        s.status.setText(f"{len(s.rows)} leads found ...")
        if len(s.rows) % 100 == 0:
            try: os.makedirs(OUT_DIR, exist_ok=True); save_excel(s.autopath, s.rows, s.fields)
            except Exception: pass
    def halt(s):
        if s.worker: s.worker.stop = True; s.status.setText("Stopping ...")
    def done(s):
        s.stopb.setEnabled(False); s.back.setEnabled(True)
        if s.rows:
            os.makedirs(OUT_DIR, exist_ok=True)
            jsave(MEM_FILE, sorted(set(jload(MEM_FILE, [])) | {r["_key"] for r in s.rows if r.get("_key")}))
            p = s.autopath; er = s.worker.errors
            try: save_excel(p, s.rows, s.fields); s.status.setText(f"Done - {len(s.rows)} leads. Saved to {p}" + (f"  ({er} search errors, see gharleads_log.txt)" if er else ""))
            except Exception as e: s.status.setText(f"Done - {len(s.rows)} leads (auto-save failed: {e})")
        else: s.status.setText("Finished - no leads found. Check your internet connection or try different stations.")
    def save_as(s):
        if not s.rows: return QMessageBox.information(s, "GharLeads", "No leads to save yet.")
        p, _ = QFileDialog.getSaveFileName(s, "Save leads", os.path.join(OUT_DIR, "gharleads_leads.xlsx"), "Excel (*.xlsx)")
        if p: save_excel(p, s.rows, s.fields); s.status.setText(f"Saved {len(s.rows)} leads to {p}")

def selftest():
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch(headless=True); pg = b.new_page(); pg.goto("https://www.google.com/maps", timeout=30000); t = pg.title(); b.close()
        msg = f"Self-test PASSED.\nThe bundled browser works and reached: {t}"
    except Exception as e: msg = f"Self-test FAILED:\n{str(e)[:500]}"
    QMessageBox.information(None, "GharLeads self-test", msg)

if __name__ == "__main__":
    app = QApplication(sys.argv)
    if "--selftest" in sys.argv: selftest(); sys.exit(0)
    w = App(); w.show(); sys.exit(app.exec())